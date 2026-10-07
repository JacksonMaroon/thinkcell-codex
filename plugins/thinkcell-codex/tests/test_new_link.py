"""New-link guards use neutral models; no proprietary native asset is bundled."""
import argparse
import base64
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import uuid
import zipfile
from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills/thinkcell-edit/scripts'
sys.path.insert(0, str(SCRIPTS / 'external_links'))
import new_link as new

DATA = dict(series_names=['S1','S2','S3','S4'], categories=['C1','C2'], series_values=[[1.,2.],[3.,4.],[5.,6.],[7.,8.]])
XML = b'''<root><version val="38732"/><CSmartGrid id="1"><m_cadvisesink/></CSmartGrid>
<CSequenceChartSE id="2"><m_ect val="0"/><m_dtable idref="3"/></CSequenceChartSE>
<CSequenceChartDataTable id="3"><m_strName>chart</m_strName><m_bstrRangeName val="think-cellChild0"/>
<m_bExternalStorage val="1"/><m_advisesink idref="0"/></CSequenceChartDataTable></root>'''

class NewLinkGuards(unittest.TestCase):
    def profile(self, xml=XML):
        with patch.object(new, 'sequence_tables', return_value=[DATA]):
            return new.profile(xml, 'chart', False)

    def test_unlinked_profile(self):
        self.assertEqual(self.profile()[1].get('id'), '3')

    def test_wrong_version_type_ownership_name_or_existing_link(self):
        for before, after in [(b'38732',b'35750'), (b'val="0"', b'val="1"'),
                              (b'm_dtable idref="3"',b'm_dtable idref="2"'),
                              (b'>chart<',b'>sibling<'), (b'm_advisesink idref="0"',b'm_advisesink idref="8"'),
                              (b'id="3"', b'id="2"')]:
            with self.subTest(after=after), self.assertRaises(RuntimeError):
                self.profile(XML.replace(before, after))

    def test_shape_and_duplicate_series_guards(self):
        for data in [dict(DATA, categories=['C1']), dict(DATA,series_names=['S1']*4), dict(DATA,series_values=[[None,2.]]*4)]:
            with patch.object(new, 'sequence_tables',return_value=[data]), self.assertRaises(RuntimeError):
                new.profile(XML,'chart',False)

    def test_fresh_native_encoding_and_collision_retry(self):
        header = bytes(range(12))
        first = uuid.UUID('00112233-4455-4677-8899-aabbccddeeff')
        second = uuid.UUID('00112233-4455-4677-8899-aabbccddee00')
        with patch.object(new.uuid,'uuid4',side_effect=[first,second]):
            guid, link, hidden = new.fresh_identity(header,{str(first)})
        self.assertEqual(guid,str(second)); self.assertTrue(link.isalnum())
        raw = base64.b32decode(hidden[12:]+'='*(-len(hidden[12:])%8))
        self.assertEqual(raw,header+second.bytes_le)

    def test_identity_exhaustion_fails(self):
        value = uuid.UUID('00112233-4455-4677-8899-aabbccddeeff')
        with patch.object(new.uuid,'uuid4',return_value=value),self.assertRaisesRegex(RuntimeError,'allocate'):
            new.fresh_identity(bytes(12),{str(value)})

    def workbook(self,path,hidden='___thinkcellTEST',reference="'ChartData'!$A$1:$C$5",values=None):
        from openpyxl.workbook.defined_name import DefinedName
        wb=Workbook(); ws=wb.active; ws.title='ChartData'
        for row in values or [['','C1','C2'],['S1',1,2],['S2',3,4],['S3',5,6],['S4',7,8]]:
            ws.append(row)
        wb.defined_names.add(DefinedName(hidden,attr_text=reference,localSheetId=0,hidden=True))
        wb.save(path); wb.close()

    def test_workbook_range_and_numeric_data_guards(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'source.xlsx';self.workbook(path)
            self.assertEqual(new.workbook_identity(path,'___thinkcellTEST','ChartData',DATA)['range'],'A1:C5')
            self.workbook(path,reference="'ChartData'!A1:C6")
            with self.assertRaisesRegex(RuntimeError,'A1:C5'):
                new.workbook_identity(path,'___thinkcellTEST','ChartData',DATA)
            self.workbook(path,values=[['','C1','C2'],['S1',99,2],['S2',3,4],['S3',5,6],['S4',7,8]])
            with self.assertRaisesRegex(RuntimeError,'numeric'):
                new.workbook_identity(path,'___thinkcellTEST','ChartData',DATA)

    def test_reordered_or_duplicate_series_rows_are_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'source.xlsx'
            for rows in [
                [['','C1','C2'],['S2',3,4],['S1',1,2],['S3',5,6],['S4',7,8]],
                [['','C1','C2'],['S1',1,2],['S1',3,4],['S3',5,6],['S4',7,8]],
            ]:
                with self.subTest(rows=rows):
                    self.workbook(path,values=rows)
                    with self.assertRaisesRegex(RuntimeError,'ordered series'):
                        new.workbook_identity(path,'___thinkcellTEST','ChartData',DATA)

    def test_numeric_matrix_is_compared_by_series_position(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'source.xlsx'
            self.workbook(path,values=[['','C1','C2'],['S1',3,4],['S2',1,2],['S3',5,6],['S4',7,8]])
            with self.assertRaisesRegex(RuntimeError,'numeric'):
                new.workbook_identity(path,'___thinkcellTEST','ChartData',DATA)

    def test_multiple_native_range_aliases_are_rejected(self):
        from openpyxl import load_workbook
        from openpyxl.workbook.defined_name import DefinedName
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'source.xlsx';self.workbook(path)
            wb=load_workbook(path)
            wb.defined_names.add(DefinedName('___thinkcellALIAS',attr_text="'ChartData'!$A$1:$C$5",localSheetId=0,hidden=True))
            wb.save(path);wb.close()
            with self.assertRaisesRegex(RuntimeError,'aliases'):
                new.workbook_identity(path,'___thinkcellTEST','ChartData',DATA)

    def test_path_alias_or_existing_output_has_no_writes(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);src=base/'src';donor=base/'donor';wb=base/'wb'
            for p in (src,donor,wb):p.write_bytes(b'sealed')
            output=base/'out';report=base/'report';workout=base/'new.xlsx'
            args=argparse.Namespace(source_presentation=src,donor_presentation=donor,source_workbook=wb,
                                    output_presentation=output,output_workbook=workout,report=report)
            args.output_workbook=src
            with self.assertRaisesRegex(RuntimeError,'distinct'):new.prepare(args)
            self.assertFalse(output.exists());self.assertFalse(report.exists())
            args.output_workbook=workout;output.write_bytes(b'existing')
            with self.assertRaisesRegex(RuntimeError,'new'):new.prepare(args)
            self.assertEqual(output.read_bytes(),b'existing');self.assertFalse(workout.exists())

    def test_invalid_profile_aborts_before_any_deliverable_write(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder); src,donor,wb=[base/n for n in ('src.pptx','donor.pptx','src.xlsx')]
            for path in (src,donor):
                with zipfile.ZipFile(path,'w') as z:z.writestr('ppt/embeddings/chart.bin',b'carrier')
            wb.write_bytes(b'workbook')
            args=argparse.Namespace(source_presentation=src,donor_presentation=donor,source_workbook=wb,
                output_presentation=base/'out.pptx',output_workbook=base/'out.xlsx',report=base/'report.json',
                source_sha256=new.rebind.sha256(src),donor_sha256=new.rebind.sha256(donor),
                source_workbook_sha256=new.rebind.sha256(wb),sheet_name='ChartData',range='A1:C5',
                target_part='ppt/embeddings/chart.bin',donor_part='ppt/embeddings/chart.bin',
                slide_number=1,donor_slide_number=1,chart_name='chart')
            with patch.object(new,'owner'),patch.object(new,'read_carrier',return_value=XML.replace(b'38732',b'35750')):
                with self.assertRaisesRegex(RuntimeError,'version'):new.prepare(args)
            self.assertEqual(sorted(p.name for p in base.iterdir()),['donor.pptx','src.pptx','src.xlsx'])

    def test_multiple_logical_owners_and_sibling_slide_are_rejected(self):
        pns=new.NS['p'];ans=new.NS['a'];rns=new.NS['r']
        def frame(identity):
            return f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{identity}"/></p:nvGraphicFramePr><a:graphic><a:graphicData><p:oleObj r:id="r1"/></a:graphicData></a:graphic></p:graphicFrame>'
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'deck.pptx'
            for frames, slide_number in [(frame(1),2),(frame(1)+frame(2),1)]:
                with zipfile.ZipFile(path,'w') as z:
                    z.writestr('slide.xml',f'<p:sld xmlns:p="{pns}" xmlns:a="{ans}" xmlns:r="{rns}">{frames}</p:sld>')
                with zipfile.ZipFile(path) as z, patch.object(new,'logical_slides',return_value=[{'part':'slide.xml'}]),patch.object(new,'relationship_map',return_value={'r1':{'resolved':'carrier.bin'}}):
                    with self.assertRaisesRegex(RuntimeError,'ownership|multiple'):new.owner(z,'carrier.bin',slide_number)

if __name__=='__main__':unittest.main()
