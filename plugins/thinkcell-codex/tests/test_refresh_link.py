"""Fail-closed request and exact native data verification contracts."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1]/'skills/thinkcell-edit/scripts'
sys.path.insert(0, str(SCRIPTS/'external_links'))
import refresh_link as refresh


def request():
    return dict(schema='thinkcell-refresh-v1', input_presentation='source.pptx', input_sha256='A'*64,
        source_workbook='source.xlsx', source_workbook_sha256='B'*64,
        output_presentation='output.pptx', output_workbook='output.xlsx', report='result.json',
        preview='preview.png', guid='guid', link_id='id', range_name='range', sheet='Sales',
        expected_range='A1:C4', slide=1, name='Chart',
        updates=[dict(cell='B3', expected=1, value=2)],
        expected_sequence=dict(series_names=['Sales'], categories=['FY25','FY26'], series_values=[[2,3]]),
        expected_datasheet=[[None,'FY25','FY26'],['Sales',2,3]])


def linked_fixture():
    return dict(fields={'auto_update':'1', 'advisesink_id':'2', 'advisesink_idref':'2'},
                xml=b'<root><CSmartGrid id="1"><m_cadvisesink><elem idref="2"/></m_cadvisesink></CSmartGrid></root>',part='linked.bin')


class RequestTests(unittest.TestCase):
    def setUp(self):
        self.ordinary_scope_patch=patch.object(refresh,'selected_ordinary_snapshot',return_value=None)
        self.ordinary_scope_patch.start();self.addCleanup(self.ordinary_scope_patch.stop)

    def test_valid_request_retains_explicit_numeric_cells(self):
        r = request()
        self.assertEqual(refresh.load_request(r), r)

    def test_malformed_cells_and_outside_range_rejected(self):
        for cell in ('A0','XFE1','A1048577','$A$1','a1','A1:B2','A1, B2','A1;Quit','C5'):
            with self.subTest(cell=cell), self.assertRaises(ValueError):
                r=request();r['updates'][0]['cell']=cell;refresh.load_request(r)

    def test_updates_fail_closed_before_mutation(self):
        for update in ([], [dict(cell='B3',value=2)], [dict(cell='B3',expected=1,value=1)],
                       [dict(cell='B3',expected=1,value=True)], [dict(cell='B3',expected=1,value=float('nan'))],
                       [dict(cell='B3',expected=1,value='=1+1')], [dict(cell='B3',expected=1,value='@SUM(A1)')]):
            with self.subTest(update=update), self.assertRaises(ValueError):
                r=request();r['updates']=update;refresh.load_request(r)
        r=request();r['updates']*=2
        with self.assertRaisesRegex(ValueError,'duplicate'):refresh.load_request(r)

    def test_invalid_range_schema_hash_and_dimensions(self):
        for key,value in [('expected_range','C4:A1'),('expected_range','A1:B2:C3'),('input_sha256','xyz'),
                          ('slide',True),('schema','old'),('expected_datasheet',[[1],[1,2]]),
                          ('expected_sequence',dict(series_names=['A'],categories=['x'],series_values=[[1,2]]))]:
            with self.subTest(key=key),self.assertRaises(ValueError):
                r=request();r[key]=value;refresh.load_request(r)
        r=request();r['unknown']=1
        with self.assertRaises(ValueError):refresh.load_request(r)

    def test_source_hash_mismatch_never_creates_candidate(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);r=request()
            for key in ('input_presentation','source_workbook','output_presentation','output_workbook','report','preview'):
                r[key]=str(root/r[key])
            Path(r['input_presentation']).write_bytes(b'sealed source')
            Path(r['source_workbook']).write_bytes(b'sealed workbook')
            manifest=root/'request.json';manifest.write_text(json.dumps(r))
            with self.assertRaisesRegex(ValueError,'hash mismatch'):refresh.prepare(manifest)
            self.assertFalse(Path(r['output_workbook']).exists())
            self.assertFalse(Path(r['output_presentation']).exists())

    def test_failure_does_not_delete_competing_writer_output(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);r=request()
            for key in ('input_presentation','source_workbook','output_presentation','output_workbook','report','preview'):
                r[key]=str(root/r[key])
            Path(r['input_presentation']).write_bytes(b'source');Path(r['source_workbook']).write_bytes(b'workbook')
            r['input_sha256']=refresh.sha256(Path(r['input_presentation']));r['source_workbook_sha256']=refresh.sha256(Path(r['source_workbook']))
            manifest=root/'request.json';manifest.write_text(json.dumps(r))
            def competing_writer(_):
                Path(r['output_presentation']).write_bytes(b'other writer owns this')
                raise RuntimeError('late failure')
            with patch.object(refresh,'discover',return_value={'part':'linked.bin'}),patch.object(refresh,'check_owner'),patch.object(refresh,'rebind_main',side_effect=competing_writer),self.assertRaisesRegex(RuntimeError,'late failure'):
                refresh.prepare(manifest)
            self.assertEqual(Path(r['output_presentation']).read_bytes(),b'other writer owns this')
            self.assertFalse(Path(r['output_workbook']).exists())

    def test_wrong_slide_rejected_before_candidate_or_office(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);r=request()
            for key in ('input_presentation','source_workbook','output_presentation','output_workbook','report','preview'):
                r[key]=str(root/r[key])
            with zipfile.ZipFile(r['input_presentation'],'w') as z:
                z.writestr('ppt/presentation.xml','<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="one"/><p:sldId id="257" r:id="two"/></p:sldIdLst></p:presentation>')
                z.writestr('ppt/_rels/presentation.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="one" Target="slides/slide1.xml"/><Relationship Id="two" Target="slides/slide2.xml"/></Relationships>')
                z.writestr('ppt/slides/slide1.xml','<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
                z.writestr('ppt/slides/slide2.xml','<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:oleObj r:id="linked"/></p:sld>')
                z.writestr('ppt/slides/_rels/slide2.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="linked" Target="../embeddings/linked.bin"/></Relationships>')
            Path(r['source_workbook']).write_bytes(b'workbook')
            r['input_sha256']=refresh.sha256(Path(r['input_presentation']));r['source_workbook_sha256']=refresh.sha256(Path(r['source_workbook']))
            manifest=root/'request.json';manifest.write_text(json.dumps(r))
            with patch.object(refresh,'discover',return_value={'part':'ppt/embeddings/linked.bin'}),self.assertRaisesRegex(ValueError,'slide ownership'):
                refresh.prepare(manifest)
            self.assertFalse(Path(r['output_workbook']).exists())
            refresh.check_owner(Path(r['input_presentation']),'ppt/embeddings/linked.bin',2)

    def test_shared_unselected_dependency_retains_exact_content(self):
        with tempfile.TemporaryDirectory() as folder:
            paths=[Path(folder)/'baseline.pptx',Path(folder)/'unchanged.pptx',Path(folder)/'corrupted.pptx']
            for index,path in enumerate(paths):
                with zipfile.ZipFile(path,'w') as z:
                    z.writestr('ppt/presentation.xml','<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="one"/><p:sldId id="257" r:id="two"/></p:sldIdLst><p:sldSz cx="1000" cy="700"/></p:presentation>')
                    z.writestr('ppt/_rels/presentation.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="one" Target="slides/slide1.xml"/><Relationship Id="two" Target="slides/slide2.xml"/></Relationships>')
                    for slide in (1,2):
                        z.writestr(f'ppt/slides/slide{slide}.xml','<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:cSld><p:spTree><p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="1" name="shared object"/></p:nvGraphicFramePr><p:oleObj r:id="shared" progId="Neutral.Test"/></p:graphicFrame></p:spTree></p:cSld></p:sld>')
                        z.writestr(f'ppt/slides/_rels/slide{slide}.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="shared" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="../embeddings/shared.bin"/></Relationships>')
                    z.writestr('ppt/embeddings/shared.bin',b'original exact dependency' if index<2 else b'corrupted shared dependency')
            self.assertEqual(refresh.sibling_snapshot(paths[0],1),refresh.sibling_snapshot(paths[1],1))
            self.assertNotEqual(refresh.sibling_snapshot(paths[0],1),refresh.sibling_snapshot(paths[2],1))
            with self.assertRaisesRegex(ValueError,'slide ownership'):
                refresh.check_owner(paths[0],'ppt/embeddings/shared.bin',1)

    def test_selected_title_change_detected_with_native_graph_unchanged(self):
        self.ordinary_scope_patch.stop()
        linked=dict(part='ppt/embeddings/selected.bin',xml=b'<root><CSequenceChartSE id="8"><m_pptseqchart idref="18"/></CSequenceChartSE><CPptSequenceChart id="18"><m_bstrShapeName>native-tag</m_bstrShapeName></CPptSequenceChart></root>')
        with tempfile.TemporaryDirectory() as folder:
            paths=[Path(folder)/'baseline.pptx',Path(folder)/'changed-title.pptx',Path(folder)/'changed-unknown.pptx']
            for index,path in enumerate(paths):
                with zipfile.ZipFile(path,'w') as z:
                    z.writestr('ppt/presentation.xml','<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="one"/></p:sldIdLst><p:sldSz cx="1000" cy="700"/></p:presentation>')
                    z.writestr('ppt/_rels/presentation.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="one" Target="slides/slide1.xml"/></Relationships>')
                    title='altered ordinary title' if index==1 else 'original ordinary title'
                    unknown='altered unknown tagged content' if index==2 else 'original unknown tagged content'
                    z.writestr('ppt/slides/slide1.xml',f'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:cSld><p:spTree><p:sp><p:nvSpPr><p:cNvPr id="1" name="Title"/></p:nvSpPr><p:txBody><a:p><a:r><a:t>{title}</a:t></a:r></a:p></p:txBody></p:sp><p:sp><p:nvSpPr><p:cNvPr id="4" name="Unknown tagged shape"/><p:nvPr><p:tags r:id="unknownTags"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>{unknown}</a:t></a:r></a:p></p:txBody></p:sp><p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="2" name="native chart"/><p:nvPr><p:tags r:id="nativeTags"/></p:nvPr></p:nvGraphicFramePr><c:chart r:id="cache"/></p:graphicFrame><p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="3" name="owned carrier"/></p:nvGraphicFramePr><p:oleObj r:id="carrier"/></p:graphicFrame></p:spTree></p:cSld></p:sld>')
                    z.writestr('ppt/slides/_rels/slide1.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="carrier" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="../embeddings/selected.bin"/><Relationship Id="cache" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="../charts/chart1.xml"/><Relationship Id="nativeTags" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags" Target="../tags/native.xml"/><Relationship Id="unknownTags" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags" Target="../tags/unknown.xml"/></Relationships>')
                    z.writestr('ppt/tags/native.xml','<tags><tag name="THINKCELLSHAPEDONOTDELETE" val="native-tag"/></tags>')
                    z.writestr('ppt/tags/unknown.xml','<tags><tag name="THINKCELLSHAPEDONOTDELETE" val="unknown-tag"/></tags>')
                    z.writestr('ppt/embeddings/selected.bin',b'same authenticated native data')
                    z.writestr('ppt/charts/chart1.xml','<chart>same exact cache</chart>')
            baseline=refresh.selected_ordinary_snapshot(paths[0],1,linked)
            for path in paths[1:]:
                self.assertNotEqual(baseline,refresh.selected_ordinary_snapshot(path,1,linked))

    def test_auto_flag_is_not_refresh_evidence(self):
        linked=linked_fixture()
        plan=dict(request=request(),identity=linked['fields'],model_ids=[['CSmartGrid','1']])
        with patch.object(refresh,'discover',return_value=linked),patch.object(refresh,'sequence_tables',return_value=[]):
            with self.assertRaisesRegex(RuntimeError,'expected sequence'):refresh.verify(plan,Path('unused'))

    def test_identity_loss_fails_even_with_updated_values(self):
        plan=dict(request=request(),identity={'guid_link':'original'},model_ids={})
        with patch.object(refresh,'discover',return_value=dict(fields={'guid_link':'changed'})):
            with self.assertRaisesRegex(RuntimeError,'identity changed'):refresh.verify(plan,Path('unused'))

    def test_model_exact_cache_stale_does_not_pass(self):
        r=request();expected=r['expected_sequence'];linked=linked_fixture()
        plan=dict(request=r,identity=linked['fields'],model_ids=[['CSmartGrid','1']])
        old_model=dict(automation_name='Chart',series_values=[[1,3]])
        audit=dict(slides=1,shared_embedding_owners=[],assertions={'strict_parity_pass':True},
                   slide_details=[dict(active_documents=[{'part':'linked.bin','model':{'sequence_tables':[old_model]}}],native_charts=[dict(series=[dict(name='Sales',categories=['FY25','FY26'],values=[1,3])])])])
        with patch.object(refresh,'discover',return_value=linked),patch.object(refresh,'sequence_tables',return_value=[dict(automation_name='Chart',**expected)]),patch.object(refresh,'inspect_presentation',return_value=audit),patch.object(refresh,'workbook_preservation',return_value=True):
            with self.assertRaisesRegex(RuntimeError,'physical chart cache'):refresh.verify(plan,Path('unused'))

    def test_changed_ownership_cannot_be_suppressed(self):
        r=request();linked=linked_fixture()
        plan=dict(request=r,identity=linked['fields'],model_ids=[['CSmartGrid','1']])
        audits=[dict(slides=1,shared_embedding_owners=['new']),dict(slides=1,shared_embedding_owners=[])]
        with patch.object(refresh,'discover',return_value=linked),patch.object(refresh,'sequence_tables',return_value=[dict(automation_name='Chart',**r['expected_sequence'])]),patch.object(refresh,'inspect_presentation',side_effect=audits),patch.object(refresh,'workbook_preservation',return_value=True):
            with self.assertRaisesRegex(RuntimeError,'ownership records changed'):refresh.verify(plan,Path('unused'))

    def test_unrelated_workbook_content_and_style_changes_fail(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as folder:
            source,target=Path(folder)/'source.xlsx',Path(folder)/'target.xlsx'
            book=Workbook();book.active.title='Sales';book.active['B3']=1;book.active['D4']='=SUM(B3,5)';book.save(source);book.close()
            for corruption in ('formula','value','style'):
                book=load_workbook(source);book['Sales']['B3']=2
                if corruption=='formula':book['Sales']['D4']='=SUM(B3,6)'
                if corruption=='value':book['Sales']['A1']='unrelated change'
                if corruption=='style':book['Sales']['D4'].number_format='0.0%'
                book.save(target);book.close()
                with self.subTest(corruption=corruption),self.assertRaisesRegex(RuntimeError,'unrelated cell/formula/name/style'):
                    refresh.workbook_preservation(source,target,[dict(cell='B3',expected=1,value=2)],'Sales')

    def test_sibling_content_change_fails_without_suppressing_gate(self):
        linked=linked_fixture();r=request();plan=dict(request=r,identity=linked['fields'],model_ids=[['CSmartGrid','1']])
        audits=dict(slides=2)
        with patch.object(refresh,'discover',return_value=linked),patch.object(refresh,'sequence_tables',return_value=[dict(automation_name='Chart',**r['expected_sequence'])]),patch.object(refresh,'inspect_presentation',return_value=audits),patch.object(refresh,'sibling_snapshot',side_effect=[{'text':'altered'},{'text':'original'}]):
            with self.assertRaisesRegex(RuntimeError,'unselected slide'):refresh.verify(plan,Path('unused'))

    def test_native_adapter_has_bounded_polling_and_scoped_cleanup(self):
        text=(SCRIPTS/'external_links/refresh_link_native.ps1').read_text()
        public=(SCRIPTS/'external_links/refresh_link.py').read_text()
        self.assertIn("OfficeOperationLock('persistent-excel-refresh')",public)
        self.assertIn('run_locked_subprocess',public)
        self.assertNotIn('Kearney',text)
        self.assertIn('$timer.Elapsed.TotalSeconds-lt$TimeoutSeconds',text)
        self.assertIn('validate-plan',text)
        self.assertIn('other_workbooks_unchanged',text)
        for forbidden in ('.Quit(','.Activate(','.Selection','UpdateBatch(','.Send('):
            self.assertNotIn(forbidden,text)


if __name__=='__main__':unittest.main()
