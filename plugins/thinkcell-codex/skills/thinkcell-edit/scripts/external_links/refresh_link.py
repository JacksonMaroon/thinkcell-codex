"""Guarded persistent Excel refresh on disposable, explicitly named copies.

prepare never opens Office. execute invokes the serialized native adapter.
Verification requires expected data, not merely an enabled refresh flag.
"""
from __future__ import annotations
import argparse
import contextlib
import io
import json
import math
import posixpath
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from portable_rebind import discover, main as rebind_main, replace_xml_stream, sha256, _path_runs, named_range
from audit_thinkcell_integrity import inspect_presentation, sequence_tables, logical_slides, relationship_map, NS, R
from office_operation_lock import OfficeOperationLock, run_locked_subprocess
from runtime import powershell, powershell_env
from extract_thinkcell_named_datasheet import extract_named_datasheet
from lxml import etree


def read_binary_grid(path):
    from pyxlsb import open_workbook
    from pyxlsb.handlers import CellHandler
    class InlineStringHandler(CellHandler):
        def read(self, reader, recid, reclen):
            column, style = reader.read_int(), reader.read_int()
            return self.cls(column, reader.read_string(), None, style)
    with open_workbook(str(path)) as book, book.get_sheet(1) as sheet:
        # pyxlsb omits inline BrtCellSt (0x0006), emitted by native think-cell.
        # Register it on this reader only, preserving the library's global map.
        sheet._reader.handlers = dict(sheet._reader.handlers)
        sheet._reader.register_handler(6, InlineStringHandler())
        return {row[0].r: {cell.c: cell.v for cell in row} for row in sheet.rows() if row}


def workbook_preservation(source, target, updates, selected_sheet):
    from openpyxl import load_workbook
    books = [load_workbook(path, data_only=False, read_only=False) for path in (source, target)]
    def xml(value):
        import copy
        node = copy.copy(value).to_tree()
        return etree.tostring(node, method='c14n')
    def sheet_snapshot(sheet):
        cells = {}
        for row in sheet:
            for cell in row:
                if cell.value is not None or cell.has_style or cell.comment or cell.hyperlink:
                    cells[cell.coordinate] = (cell.value, cell.data_type, cell.number_format,
                        xml(cell.font), xml(cell.fill), xml(cell.border), xml(cell.alignment), xml(cell.protection),
                        None if cell.comment is None else (cell.comment.text, cell.comment.author),
                        None if cell.hyperlink is None else xml(cell.hyperlink))
        return {'cells': cells, 'names': sorted(xml(v) for v in sheet.defined_names.values()),
                'merged': str(sheet.merged_cells), 'rows': {k:xml(v) for k,v in sheet.row_dimensions.items()},
                'columns': {k:xml(v) for k,v in sheet.column_dimensions.items()},
                'protection':xml(sheet.protection), 'state':sheet.sheet_state}
    try:
        before, after = books
        if before.sheetnames != after.sheetnames or sorted(xml(v) for v in before.defined_names.values()) != sorted(xml(v) for v in after.defined_names.values()):
            raise RuntimeError('workbook sheets or defined names changed')
        for name in before.sheetnames:
            left, right = sheet_snapshot(before[name]), sheet_snapshot(after[name])
            if name == selected_sheet:
                for update in updates:
                    cell = update['cell']
                    expected = list(left['cells'][cell])
                    expected[0] = update['value']; expected[1] = 's' if isinstance(update['value'], str) else 'n'
                    left['cells'][cell] = tuple(expected)
            if left != right:
                raise RuntimeError('workbook unrelated cell/formula/name/style semantics changed: '+name)
    finally:
        for book in books:
            book.close()
    return True


def sibling_snapshot(path, slide):
    from assemble_native_slides import physical_snapshot, dependencies, inspect_shapes
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as before, zipfile.ZipFile(buffer, 'w') as after:
        for item in before.infolist():
            data = before.read(item.filename)
            if item.filename.endswith('.xml'):
                root = etree.fromstring(data)
                for node in root.iter('{'+NS['p']+'}oleObj'):
                    if node.get('progId') == 'TCLayout.ActiveDocument.1' and node.get('name') == 'think-cell Slide':
                        node.set('name', 'thinkcell Slide')
                data = etree.tostring(root)
            after.writestr(item, data)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as package:
        slides = logical_slides(package)
        if not 1 <= slide <= len(slides):
            raise RuntimeError('sibling slide is outside presentation')
        part = slides[slide-1]['part']
        inspect_shapes(package.read(part))
        dimensions = etree.fromstring(package.read('ppt/presentation.xml')).find('p:sldSz', NS)
        if dimensions is None:
            raise RuntimeError('presentation dimensions missing')
        notes = [relationship['resolved'] for relationship in relationship_map(package, part).values() if relationship['type'].endswith('/notesSlide')]
        if len(notes) > 1:
            raise RuntimeError('multiple sibling notes dependencies')
        # The immutable readback scope is this slide's complete physical and
        # dependency closure. Global inventory rejects pre-existing shared
        # unselected carriers even when their complete stream content and exact
        # ownership records are preserved. Ownership remains separately gated.
        return {'dimensions':tuple(sorted(dimensions.attrib.items())),
                'physical':physical_snapshot(package, part),
                'notes':physical_snapshot(package, notes[0]) if notes else None,
                'dependencies':dependencies(package, part), 'slide_count':len(slides),
                'slide_ids':[entry['id'] for entry in slides]}


def selected_ordinary_snapshot(path, slide_number, linked):
    """Exclude only shapes authenticated by the selected native ownership graph."""
    root = etree.fromstring(linked['xml'])
    ids = {node.get('id'):node for node in root if node.get('id')}
    charts = root.findall('CSequenceChartSE')
    if len(charts) != 1:
        raise RuntimeError('selected ordinary-content guard requires one native chart owner')
    pending, reachable = [charts[0].get('id')], set()
    while pending:
        identity = pending.pop()
        if identity == '0' or identity in reachable:
            continue
        if identity not in ids:
            raise RuntimeError('native shape ownership graph has a dangling reference')
        reachable.add(identity)
        pending.extend(node.get('idref') for node in ids[identity].iter() if node.get('idref'))
    owned_tags = {node.text for identity in reachable for node in ids[identity].iter('m_bstrShapeName') if node.text}
    with zipfile.ZipFile(path) as package:
        parts = {item.filename:package.read(item.filename) for item in package.infolist()}
        slides = logical_slides(package);part = slides[slide_number-1]['part']
        rels = relationship_map(package, part)
        tree = etree.fromstring(parts[part]); shapes = tree.find('p:cSld/p:spTree', NS)
        if shapes is None:
            raise RuntimeError('selected ordinary shape tree missing')
        removed_refs = set();removed = 0;owned_chart_frames = 0;tag_owners = set()
        for shape in list(shapes.iter()):
            if etree.QName(shape).namespace != NS['p'] or etree.QName(shape).localname not in {'sp','cxnSp','graphicFrame','pic'}:
                continue
            tags = []
            for tag_reference in shape.findall('.//p:tags', NS):
                relationship = rels.get(tag_reference.get(R+'id'))
                if not relationship or relationship['resolved'] not in parts:
                    raise RuntimeError('selected shape tag relationship does not resolve')
                tag_root = etree.fromstring(parts[relationship['resolved']])
                tags.extend(node.get('val') for node in tag_root if node.get('name','').upper() == 'THINKCELLSHAPEDONOTDELETE')
            native_tag_owner = bool(tags) and len(tags) == 1 and tags[0] in owned_tags
            if native_tag_owner:
                if tags[0] in tag_owners:
                    raise RuntimeError('native shape tag has ambiguous physical ownership')
                tag_owners.add(tags[0])
            linked_ole_owner = any(rels.get(node.get(R+'id'),{}).get('resolved') == linked['part'] for node in shape.findall('.//p:oleObj', NS))
            if native_tag_owner or linked_ole_owner:
                if shape.find('.//c:chart', NS) is not None:
                    owned_chart_frames += 1
                removed_refs.update(value for node in shape.iter() for key,value in node.attrib.items() if key.startswith(R))
                shape.getparent().remove(shape);removed += 1
        if not removed or owned_chart_frames != 1:
            raise RuntimeError('exactly one authenticated native chart frame required for ordinary-content scope')
        retained_refs = {value for node in tree.iter() for key,value in node.attrib.items() if key.startswith(R)}
        relpart = posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')
        relationship_root = etree.fromstring(parts[relpart])
        for relationship in list(relationship_root):
            if relationship.get('Id') in removed_refs - retained_refs:
                relationship_root.remove(relationship)
        parts[part] = etree.tostring(tree);parts[relpart] = etree.tostring(relationship_root)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer,'w') as output:
        for name,data in parts.items():
            output.writestr(name,data)
    with tempfile.TemporaryDirectory() as folder:
        scoped = Path(folder)/'ordinary.pptx';scoped.write_bytes(buffer.getvalue())
        return sibling_snapshot(scoped, slide_number)


def coordinate(cell):
    if not isinstance(cell, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', cell):
        raise ValueError('cell must be a canonical single A1 address')
    letters, digits = re.fullmatch(r'([A-Z]+)([0-9]+)', cell).groups()
    col = 0
    for letter in letters:
        col = col * 26 + ord(letter) - 64
    row = int(digits)
    if col > 16384 or row > 1048576:
        raise ValueError('cell exceeds Excel bounds')
    return row, col


def scalar(value):
    if value is None or isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ValueError('only strings and finite numbers are supported')
    if isinstance(value, str) and (not value or value.startswith(('=', '+', '-', '@')) or any(ord(c) < 32 for c in value)):
        raise ValueError('formula-like, empty or control-character text is unsupported')
    if isinstance(value, (float, int)) and not math.isfinite(value):
        raise ValueError('finite numbers required')
    return value


def load_request(path):
    request = path if isinstance(path, dict) else json.loads(Path(path).read_text(encoding='utf-8-sig'))
    required = {'schema', 'input_presentation', 'input_sha256', 'source_workbook',
                'source_workbook_sha256', 'output_presentation', 'output_workbook',
                'report', 'preview', 'guid', 'link_id', 'range_name', 'sheet',
                'expected_range', 'slide', 'name', 'updates', 'expected_sequence', 'expected_datasheet'}
    if not isinstance(request, dict) or set(request) != required or request['schema'] != 'thinkcell-refresh-v1':
        raise ValueError('exact thinkcell-refresh-v1 request fields required')
    for key in required - {'updates', 'expected_sequence', 'expected_datasheet', 'slide', 'name'}:
        if not isinstance(request[key], str) or not request[key].strip():
            raise ValueError(f'{key} must be a nonempty string')
    if not isinstance(request['name'], str):
        raise ValueError('exact automation name string required; empty selects an unnamed linked table')
    for key in ('input_sha256', 'source_workbook_sha256'):
        if not re.fullmatch('[a-fA-F0-9]{64}', request[key]):
            raise ValueError('64-digit SHA-256 required')
    if type(request['slide']) is not int or request['slide'] < 1:
        raise ValueError('positive integer slide required')
    bounds = request['expected_range'].split(':')
    if len(bounds) > 2:
        raise ValueError('one contiguous canonical range required')
    start, end = coordinate(bounds[0]), coordinate(bounds[-1])
    if any(a > b for a, b in zip(start, end)):
        raise ValueError('range endpoints reversed')
    updates = request['updates']
    if not isinstance(updates, list) or not updates:
        raise ValueError('explicit nonempty updates required')
    seen = set()
    for update in updates:
        if not isinstance(update, dict) or set(update) != {'cell', 'expected', 'value'}:
            raise ValueError('each update requires cell, expected and value')
        row, col = coordinate(update['cell'])
        if update['cell'] in seen or not(start[0] <= row <= end[0] and start[1] <= col <= end[1]):
            raise ValueError('duplicate or out-of-range update')
        seen.add(update['cell'])
        scalar(update['expected']); scalar(update['value'])
        if update['expected'] == update['value']:
            raise ValueError('no-op updates rejected')
    expected = request['expected_sequence']
    if not isinstance(expected, dict) or set(expected) != {'series_names', 'categories', 'series_values'}:
        raise ValueError('complete expected sequence required')
    names, categories, values = (expected[k] for k in ('series_names', 'categories', 'series_values'))
    if not all(isinstance(v, list) and v for v in (names, categories, values)) or len(values) != len(names):
        raise ValueError('sequence dimensions invalid')
    for name in names + categories:
        scalar(name)
    for series in values:
        if not isinstance(series, list) or len(series) != len(categories):
            raise ValueError('sequence dimensions invalid')
        for value in series:
            if type(value) not in (int, float):
                raise ValueError('numeric series required')
            scalar(value)
    grid = request['expected_datasheet']
    if not isinstance(grid, list) or not grid or not all(isinstance(row, list) and len(row) == len(grid[0]) and row for row in grid):
        raise ValueError('rectangular expected datasheet required')
    for row in grid:
        for value in row:
            if value is not None:
                scalar(value)
    return request


def check_owner(presentation, carrier, slide_number):
    owners = []
    with zipfile.ZipFile(presentation) as package:
        for index, slide in enumerate(logical_slides(package), 1):
            rels = relationship_map(package, slide['part'])
            tree = etree.fromstring(package.read(slide['part']))
            if any(rels.get(node.get(R+'id'), {}).get('resolved') == carrier for node in tree.xpath('.//p:oleObj', namespaces=NS)):
                owners.append(index)
    if owners != [slide_number]:
        raise ValueError(f'selected linked carrier slide ownership {owners} differs from requested slide {slide_number}')


def prepare(request_path):
    r = load_request(request_path)
    paths = [Path(r[key]).resolve() for key in ('input_presentation', 'source_workbook', 'output_presentation', 'output_workbook', 'report', 'preview')]
    if len(set(paths + [Path(request_path).resolve()])) != 7:
        raise ValueError('all input, output and request paths must differ')
    source, workbook, output, target, report, preview = paths
    if source.suffix.lower() != '.pptx' or workbook.suffix.lower() != '.xlsx' or output.suffix.lower() != '.pptx' or target.suffix.lower() != '.xlsx':
        raise ValueError('PPTX and nonmacro XLSX required')
    if any(path.exists() for path in paths[2:]):
        raise ValueError('fresh output, workbook, report and preview paths required')
    if sha256(source) != r['input_sha256'].upper() or sha256(workbook) != r['source_workbook_sha256'].upper():
        raise ValueError('sealed input hash mismatch')
    selected_source = discover(source, r['guid'], r['link_id'], r['range_name'])
    check_owner(source, selected_source['part'], r['slide'])
    for path in paths[2:]:
        path.parent.mkdir(parents=True, exist_ok=True)
    created = []
    try:
        with target.open('xb') as handle:
            created.append(target)
            handle.write(workbook.read_bytes())
        with tempfile.TemporaryDirectory() as folder:
            staged, evidence = Path(folder)/'rebound.pptx', Path(folder)/'rebind.json'
            argv = ['--input-presentation', str(source), '--input-sha256', r['input_sha256'], '--source-workbook', str(workbook), '--source-workbook-sha256', r['source_workbook_sha256'], '--target-workbook', str(target), '--target-workbook-sha256', sha256(target), '--output-presentation', str(staged), '--report', str(evidence)]
            for key in ('guid', 'link_id', 'range_name', 'expected_range'):
                argv.extend(['--'+key.replace('_', '-'), r[key]])
            with contextlib.redirect_stdout(io.StringIO()):
                rebind_main(argv)
            rebind_report = json.loads(evidence.read_text())
            if rebind_report['guards']['named_range_compatibility']['target']['sheet'] != r['sheet']:
                raise ValueError('request sheet differs from sealed native link')
            selected = discover(staged, r['guid'], r['link_id'], r['range_name'])
            xml = selected['xml']
            root = ET.fromstring(xml)
            flags = list(root.iter('m_bAutoUpdate'))
            if len(flags) != 1 or flags[0].get('val') not in ('0', '1'):
                raise ValueError('unsupported auto-update flag')
            changed = re.sub(rb'(<m_bAutoUpdate\b[^>]*\bval=")[01](")', rb'\g<1>1\2', xml)
            if selected['fields']['auto_update'] == '0' and changed == xml:
                raise ValueError('unsupported auto-update serialization')
            raw = replace_xml_stream(selected['raw'], changed)
            with output.open('xb') as handle:
                created.append(output)
                with zipfile.ZipFile(staged) as before, zipfile.ZipFile(handle, 'w') as after:
                    for item in before.infolist():
                        after.writestr(item, raw if item.filename == selected['part'] else before.read(item.filename))
            after = discover(output, r['guid'], r['link_id'], r['range_name'])
            if after['fields']['auto_update'] != '1':
                raise RuntimeError('auto-update did not survive readback')
            identities = [[node.tag, node.get('id')] for node in ET.fromstring(after['xml']).iter() if node.tag in ('CSequenceChartDataTable', 'CSmartGrid')]
            plan = {'schema': 'thinkcell-refresh-plan-v1', 'request': r, 'request_sha256': sha256(Path(request_path)), 'prepared_presentation_sha256': sha256(output), 'prepared_workbook_sha256': sha256(target), 'identity': after['fields'], 'model_ids': identities, 'rebind': json.loads(evidence.read_text()), 'status': 'PREPARED_REQUIRES_NATIVE_EXECUTION'}
            return plan
    except BaseException:
        for path in created:
            if path.exists():
                path.unlink()
        raise


def verify(plan, presentation):
    r = plan['request']
    linked = discover(Path(presentation), r['guid'], r['link_id'], r['range_name'])
    persistent_keys = set(plan['identity']) - {'advisesink_id', 'advisesink_idref'}
    if any(linked['fields'].get(key) != plan['identity'][key] for key in persistent_keys):
        raise RuntimeError('persistent link identity changed')
    root = ET.fromstring(linked['xml'])
    grids = list(root.iter('CSmartGrid'))
    registrations = [] if len(grids) != 1 else list(grids[0].findall('./m_cadvisesink/elem'))
    if len(registrations) != 1 or registrations[0].get('idref') != linked['fields']['advisesink_id']:
        raise RuntimeError('native adviser table/grid registration does not resolve')
    ids = [[node.tag, node.get('id')] for node in root.iter() if node.tag in ('CSequenceChartDataTable', 'CSmartGrid')]
    if ids != plan['model_ids']:
        raise RuntimeError('native table/grid identity changed')
    tables = [t for t in sequence_tables(etree.fromstring(linked['xml'])) if t['automation_name'] == r['name']]
    if len(tables) != 1 or any(tables[0][key] != value for key, value in r['expected_sequence'].items()):
        raise RuntimeError('native model has not reached expected sequence')
    audit = inspect_presentation(Path(presentation), True)
    baseline = inspect_presentation(Path(r['input_presentation']), True)
    source_linked = discover(Path(r['input_presentation']), r['guid'], r['link_id'], r['range_name'])
    if selected_ordinary_snapshot(Path(presentation),r['slide'],linked) != selected_ordinary_snapshot(Path(r['input_presentation']),r['slide'],source_linked):
        raise RuntimeError('selected slide ordinary/unknown shape content or dependencies changed')
    if audit['slides'] != baseline['slides']:
        raise RuntimeError('presentation slide count changed')
    for slide_number in range(1, audit['slides'] + 1):
        if slide_number != r['slide'] and sibling_snapshot(Path(presentation), slide_number) != sibling_snapshot(Path(r['input_presentation']), slide_number):
            raise RuntimeError('unselected slide content/dependencies changed: '+str(slide_number))
    workbook_preservation(Path(r['source_workbook']), Path(r['output_workbook']), r['updates'], r['sheet'])
    if audit['shared_embedding_owners'] != baseline['shared_embedding_owners']:
        raise RuntimeError('embedding ownership records changed')
    failed = [key for key, value in audit['assertions'].items() if not value and key != 'embedded_parts_have_unique_owners']
    if failed:
        raise RuntimeError('strict native integrity checks failed: ' + ', '.join(failed))
    slide = audit['slide_details'][r['slide']-1]
    if linked['part'] not in [item['part'] for item in slide['active_documents']]:
        raise RuntimeError('selected linked carrier is not owned by requested slide')
    if len(slide['native_charts']) != 1:
        raise RuntimeError('refresh requires exactly one physical native chart on selected slide')
    series = slide['native_charts'][0]['series']
    expected = r['expected_sequence']
    before_slide = baseline['slide_details'][r['slide']-1]
    before_models = [table for active in before_slide['active_documents'] if active['part'] == linked['part'] for table in active['model']['sequence_tables'] if table['automation_name'] == r['name']]
    if len(before_models) != 1 or len(before_slide['native_charts']) != 1:
        raise RuntimeError('baseline must have one exact model/cache pair')
    before_series = before_slide['native_charts'][0]['series']
    old_values = before_models[0]['series_values']
    cached_old_values = [s['values'] for s in before_series]
    if cached_old_values == old_values:
        order = list(range(len(old_values)))
    elif cached_old_values == list(reversed(old_values)):
        order = list(reversed(range(len(old_values))))
    else:
        raise RuntimeError('unsupported baseline physical series ordering')
    if [s['values'] for s in series] != [expected['series_values'][i] for i in order]:
        raise RuntimeError('physical chart cache differs from exact expected data')
    for position, index in enumerate(order):
        if series[position]['name'] != (expected['series_names'][index] if before_series[position]['name'] is not None else None):
            raise RuntimeError('physical cache series label changed')
        categories = expected['categories'] if before_series[position]['categories'] else []
        if series[position]['categories'] != categories:
            raise RuntimeError('physical cache category representation changed')
    data, kind, evidence = extract_named_datasheet(Path(presentation), r['slide'], r['name'])
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder)/('datasheet.xls' if kind == 'legacy_biff_cfb' else 'datasheet.xlsb')
        path.write_bytes(data)
        if kind == 'legacy_biff_cfb':
            import xlrd
            sheet = xlrd.open_workbook(path).sheet_by_index(0)
            grid = [[sheet.cell_value(row, col) if row < sheet.nrows and col < sheet.ncols else '' for col in range(len(r['expected_datasheet'][0]))] for row in range(len(r['expected_datasheet']))]
        else:
            rows = read_binary_grid(path)
            grid = [[rows.get(row, {}).get(col) for col in range(len(r['expected_datasheet'][0]))] for row in range(len(r['expected_datasheet']))]
    normal = lambda value: None if value == '' else value
    if [[normal(v) for v in row] for row in grid] != r['expected_datasheet']:
        raise RuntimeError('embedded datasheet differs from expected grid')
    return {'link_identity_preserved': True, 'adviser_registration_resolves': True,
            'adviser_local_id_before': plan['identity']['advisesink_id'],
            'adviser_local_id_after': linked['fields']['advisesink_id'],
            'model_exact': True, 'physical_cache_exact': True, 'physical_cache_series_order': order,
            'physical_cache_names_present': any(s['name'] is not None for s in series),
            'physical_cache_categories_present': any(s['categories'] for s in series),
            'datasheet_exact': True, 'datasheet': evidence, 'unselected_slides_preserved': True,
            'selected_slide_ordinary_content_dependencies_preserved': True,
            'unrelated_workbook_cells_formulas_names_styles_preserved': True}


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest='mode', required=True)
    prep = sub.add_parser('prepare'); prep.add_argument('--request', type=Path, required=True); prep.add_argument('--plan', type=Path, required=True)
    check = sub.add_parser('verify'); check.add_argument('--plan', type=Path, required=True); check.add_argument('--presentation', type=Path, required=True)
    execute = sub.add_parser('execute'); execute.add_argument('--plan', type=Path, required=True)
    validate = sub.add_parser('validate-plan'); validate.add_argument('--plan', type=Path, required=True)
    a = p.parse_args(argv)
    if a.mode == 'prepare':
        if a.plan.exists():
            raise ValueError('fresh plan required')
        request = load_request(a.request)
        if a.plan.resolve() in {Path(request[k]).resolve() for k in ('input_presentation', 'source_workbook', 'output_presentation', 'output_workbook', 'report', 'preview')} or a.plan.resolve() == a.request.resolve():
            raise ValueError('plan must differ from every input/output')
        result = prepare(a.request)
        a.plan.parent.mkdir(parents=True, exist_ok=True)
        with a.plan.open('x', encoding='utf-8') as handle:
            handle.write(json.dumps(result, indent=2))
    elif a.mode == 'validate-plan':
        plan = json.loads(a.plan.read_text(encoding='utf-8-sig'))
        request = load_request(plan['request'])
        paths = [Path(request[key]).resolve() for key in ('input_presentation', 'source_workbook', 'output_presentation', 'output_workbook', 'report', 'preview')]
        if len(set(paths + [a.plan.resolve()])) != 7:
            raise ValueError('plan/input/output aliases rejected')
        for key, digest in (('input_presentation', request['input_sha256']), ('source_workbook', request['source_workbook_sha256']), ('output_presentation', plan['prepared_presentation_sha256']), ('output_workbook', plan['prepared_workbook_sha256'])):
            if sha256(Path(request[key])) != digest.upper():
                raise ValueError('plan sealed hash mismatch')
        linked = discover(paths[2], request['guid'], request['link_id'], request['range_name'])
        if linked['fields'] != plan['identity'] or linked['fields']['auto_update'] != '1':
            raise ValueError('prepared persistent identity mismatch')
        check_owner(paths[2], linked['part'], request['slide'])
        if [Path(item[2]).resolve() for item in _path_runs(linked['payload'])] != [paths[3]]:
            raise ValueError('prepared link must point to exact task-owned workbook')
        moniker_sheet, moniker_name = linked['moniker_text'].split(';')[-1].split('!', 1)
        binding = named_range(paths[3], moniker_name, moniker_sheet)
        if binding['range'] != request['expected_range'] or binding['sheet'] != request['sheet'] or binding != plan['rebind']['guards']['named_range_compatibility']['target']:
            raise ValueError('prepared workbook named range differs from plan')
        result = {'plan_valid': True}
    elif a.mode == 'verify':
        result = verify(json.loads(a.plan.read_text(encoding='utf-8-sig')), a.presentation)
    else:
        helper = Path(__file__).with_name('refresh_link_native.ps1')
        with OfficeOperationLock('persistent-excel-refresh'):
            completed = run_locked_subprocess([powershell(), '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'RemoteSigned', '-File', str(helper), '-Plan', str(a.plan.resolve()), '-Python', sys.executable], operation='persistent-excel-refresh-native', timeout_seconds=180, env=powershell_env(), creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        return completed.returncode
    print(json.dumps(result, indent=2)); return 0


if __name__ == '__main__':
    raise SystemExit(main())
